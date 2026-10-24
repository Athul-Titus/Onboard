import io
import pytest
from fastapi.testclient import TestClient
import docx
import pypdf
from backend.app.main import app

client = TestClient(app)

def test_parse_resume_txt():
    content = "Jane Doe\nStaff Engineer\nSkills: Python, Go, Kubernetes\nExperience: 8 years"
    files = {"file": ("resume.txt", io.BytesIO(content.encode("utf-8")), "text/plain")}
    res = client.post("/parse-resume", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["filename"] == "resume.txt"
    assert "Staff Engineer" in data["text"]
    assert data["char_count"] > 0

def test_parse_resume_docx():
    doc = docx.Document()
    doc.add_heading("Alex Rivera", 0)
    doc.add_paragraph("DevOps Engineer with expertise in Terraform, Docker, and AWS.")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    files = {"file": ("alex_resume.docx", buffer, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/parse-resume", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["filename"] == "alex_resume.docx"
    assert "DevOps Engineer" in data["text"]

def test_parse_resume_empty_file():
    files = {"file": ("empty.txt", io.BytesIO(b""), "text/plain")}
    res = client.post("/parse-resume", files=files)
    assert res.status_code == 400
    assert "empty" in res.json()["detail"].lower()

def test_parse_resume_corrupt_docx():
    files = {"file": ("broken.docx", io.BytesIO(b"not a valid zip or docx content"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/parse-resume", files=files)
    assert res.status_code == 400
    assert "corrupt" in res.json()["detail"].lower() or "unreadable" in res.json()["detail"].lower()

def test_parse_resume_unsupported_format():
    files = {"file": ("malware.exe", io.BytesIO(b"some binary code"), "application/octet-stream")}
    res = client.post("/parse-resume", files=files)
    assert res.status_code == 400
    assert "unsupported" in res.json()["detail"].lower()
